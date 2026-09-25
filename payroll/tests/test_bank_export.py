from decimal import Decimal
import datetime
import io
import csv
from openpyxl import load_workbook
from django.utils import timezone
from rest_framework.test import APITestCase
from rest_framework import status

from core.models import Organization, OrgSettings, OrganizationModule
from users.models import Employee, Role
from attendance.models import AttendancePeriod, AttendancePeriodEmployeeSnapshot
from payroll.models import (
    SalaryComponent,
    EmployeeSalaryStructure,
    EmployeeSalaryComponent,
    PayrollPeriod,
    PayrollEmployeeSnapshot,
)
from payroll.services import calculate_payroll_period, finalize_payroll_period
from payroll.exporters import BankPaymentExporter


class BankPaymentExportTests(APITestCase):
    def setUp(self):
        # Setup Org A with corporate bank settings
        self.settings_a = OrgSettings.objects.create(
            corporate_bank_name="HDFC Bank",
            corporate_account_number="998877665544",
            corporate_ifsc_code="HDFC0001234",
            corporate_bank_branch="Main Branch"
        )
        self.org_a = Organization.objects.create(name="Org A", subdomain="org-a", settings=self.settings_a)
        OrganizationModule.objects.create(organization=self.org_a, module_id='payroll', enabled=True)

        # Setup Org B
        self.settings_b = OrgSettings.objects.create(
            corporate_bank_name="ICICI Bank",
            corporate_account_number="112233445566",
            corporate_ifsc_code="ICIC0005678",
        )
        self.org_b = Organization.objects.create(name="Org B", subdomain="org-b", settings=self.settings_b)
        OrganizationModule.objects.create(organization=self.org_b, module_id='payroll', enabled=True)

        admin_role = Role.objects.filter(slug='admin').first()
        emp_role = Role.objects.filter(slug='employee').first()

        # Admin Org A
        self.admin_a = Employee.objects.create_user(
            email="admin_a@test.com", password="password123",
            organization=self.org_a, role=admin_role, is_staff=True, isSuperAdmin=True
        )

        # Admin Org B
        self.admin_b = Employee.objects.create_user(
            email="admin_b@test.com", password="password123",
            organization=self.org_b, role=admin_role, is_staff=True, isSuperAdmin=True
        )

        # Standard employee without payroll permissions in Org A
        self.viewer_a = Employee.objects.create_user(
            email="viewer_a@test.com", password="password123",
            organization=self.org_a, role=emp_role, first_name="Viewer", last_name="User"
        )

        # Employee 1 in Org A (Fully configured bank details)
        self.emp1_a = Employee.objects.create_user(
            email="emp1_a@test.com", password="password123",
            organization=self.org_a, role=emp_role, first_name="Alice", last_name="Smith",
            employee_code="EMP-0001",
            bank_name="HDFC Bank",
            account_number="50100123456789",
            ifsc_code="HDFC0000001",
            account_holder_name="Alice Smith"
        )

        # Employee 2 in Org A (Missing account number)
        self.emp2_a = Employee.objects.create_user(
            email="emp2_a@test.com", password="password123",
            organization=self.org_a, role=emp_role, first_name="Bob", last_name="Jones",
            employee_code="EMP-0002",
            bank_name="SBI",
            account_number="",  # Missing
            ifsc_code="SBIN0001111",
            account_holder_name="Bob Jones"
        )

        # Employee 3 in Org A (Missing IFSC code)
        self.emp3_a = Employee.objects.create_user(
            email="emp3_a@test.com", password="password123",
            organization=self.org_a, role=emp_role, first_name="Charlie", last_name="Brown",
            employee_code="EMP-0003",
            bank_name="Axis Bank",
            account_number="91234567890123",
            ifsc_code="",  # Missing
            account_holder_name="Charlie Brown"
        )

        # Salary Structure for Org A
        basic = SalaryComponent.objects.create(
            organization=self.org_a, name="Basic", code="BASIC",
            component_type="Earning", is_proratable=True
        )

        for emp, gross in [
            (self.emp1_a, Decimal("50000.00")),
            (self.emp2_a, Decimal("60000.00")),
            (self.emp3_a, Decimal("70000.00")),
        ]:
            struct = EmployeeSalaryStructure.objects.create(
                organization=self.org_a, employee=emp,
                effective_from=datetime.date(2026, 1, 1),
                compensation_type="MONTHLY", gross_salary=gross,
                base_net_salary=gross, is_active=True
            )
            EmployeeSalaryComponent.objects.create(salary_structure=struct, salary_component=basic, amount=gross)

        self.client.force_authenticate(user=self.admin_a)

    def _setup_attendance_and_payroll(self, finalize=True):
        att, _ = AttendancePeriod.objects.get_or_create(
            organization=self.org_a, year=2026, month=8,
            defaults={'status': 'Finalized', 'current_revision': 1, 'finalized_at': timezone.now(), 'finalized_by': self.admin_a}
        )
        if att.status != 'Finalized':
            att.status = 'Finalized'
            att.save()

        for emp in [self.emp1_a, self.emp2_a, self.emp3_a]:
            AttendancePeriodEmployeeSnapshot.objects.get_or_create(
                attendance_period=att, revision=1, employee=emp,
                defaults={
                    'is_current': True,
                    'employee_name': emp.get_full_name() or emp.email,
                    'working_days': 22,
                    'payable_attendance_units': Decimal('22.00'),
                    'present_days': Decimal('22.00'),
                    'paid_leave_days': Decimal('0.00'),
                    'unpaid_leave_days': Decimal('0.00'),
                    'absent_days': 0
                }
            )

        period = calculate_payroll_period(self.org_a, 2026, 8, self.admin_a)
        if finalize:
            period = finalize_payroll_period(self.org_a, 2026, 8, self.admin_a)
        return period

    def test_01_export_requires_finalized_payroll(self):
        """Cannot export if payroll period is in Draft or Calculated status."""
        period = self._setup_attendance_and_payroll(finalize=False)
        self.assertEqual(period.status, 'Calculated')

        url = f"/api/payroll/periods/2026/8/export-bank-file/"

        # GET pre-check should fail
        res_get = self.client.get(url)
        self.assertEqual(res_get.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertIn("non-finalized", res_get.data.get("error", ""))

        # POST file download should fail
        res_post = self.client.post(url, {"template": "HDFC", "format": "XLSX"})
        self.assertEqual(res_post.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertIn("non-finalized", res_post.data.get("error", ""))

    def test_02_export_precheck_summary(self):
        """Pre-check endpoint returns counts of payable and unpayable employees."""
        self._setup_attendance_and_payroll(finalize=True)

        url = f"/api/payroll/periods/2026/8/export-bank-file/"
        res = self.client.get(url, {"template": "HDFC"})
        self.assertEqual(res.status_code, status.HTTP_200_OK)

        data = res.data
        self.assertEqual(data['period_year'], 2026)
        self.assertEqual(data['period_month'], 8)
        self.assertEqual(data['period_status'], 'Finalized')
        self.assertEqual(data['template'], 'HDFC')
        self.assertEqual(data['total_payable_count'], 1)  # Only emp1_a has valid details
        self.assertEqual(Decimal(data['total_payable_amount']), Decimal('50000.00'))
        self.assertEqual(data['unpayable_count'], 2)  # emp2_a and emp3_a
        self.assertEqual(len(data['unpayable_employees']), 2)

        # Verify unpayable missing fields
        unpayable_map = {u['employee_code']: u['missing_fields'] for u in data['unpayable_employees']}
        self.assertIn('account_number', unpayable_map.get('EMP-0002', []))
        self.assertIn('ifsc_code', unpayable_map.get('EMP-0003', []))

    def test_03_correct_columns_generated_for_hdfc_template(self):
        """HDFC template produces the exact specified columns in XLSX and CSV."""
        self._setup_attendance_and_payroll(finalize=True)
        url = f"/api/payroll/periods/2026/8/export-bank-file/"

        expected_columns = [
            'Beneficiary Name',
            'Account Number',
            'IFSC Code',
            'Amount',
            'Payment Mode',
            'Remarks',
        ]

        # 1. CSV test
        res_csv = self.client.post(url, {"template": "HDFC", "format": "CSV"})
        self.assertEqual(res_csv.status_code, status.HTTP_200_OK)
        self.assertEqual(res_csv['Content-Type'], 'text/csv; charset=utf-8')
        self.assertIn('Bank_Payment_HDFC_2026_08.csv', res_csv['Content-Disposition'])

        csv_content = res_csv.content.decode('utf-8-sig')
        reader = list(csv.reader(io.StringIO(csv_content)))
        self.assertEqual(reader[0], expected_columns)
        self.assertEqual(len(reader), 2)  # Header + 1 payable employee
        self.assertEqual(reader[1][0], "Alice Smith")
        self.assertEqual(reader[1][1], "50100123456789")
        self.assertEqual(reader[1][2], "HDFC0000001")
        self.assertEqual(reader[1][3], "50000.00")
        self.assertEqual(reader[1][4], "NEFT")

        # 2. XLSX test
        res_xlsx = self.client.post(url, {"template": "HDFC", "format": "XLSX"})
        self.assertEqual(res_xlsx.status_code, status.HTTP_200_OK)
        self.assertIn('application/vnd.openxmlformats-officedocument.spreadsheetml.sheet', res_xlsx['Content-Type'])
        self.assertIn('Bank_Payment_HDFC_2026_08.xlsx', res_xlsx['Content-Disposition'])

        wb = load_workbook(io.BytesIO(res_xlsx.content))
        ws = wb.active
        headers = [cell.value for cell in ws[1]]
        self.assertEqual(headers, expected_columns)

        # Check data row
        row_values = [cell.value for cell in ws[2]]
        self.assertEqual(row_values[0], "Alice Smith")
        self.assertEqual(str(row_values[1]), "50100123456789")
        self.assertEqual(row_values[2], "HDFC0000001")
        self.assertEqual(float(row_values[3]), 50000.0)
        # Ensure account number is formatted as string in Excel cell
        self.assertEqual(ws.cell(row=2, column=2).data_type, 's')

    def test_04_correct_columns_generated_for_generic_template(self):
        """GENERIC_NEFT template generates correct columns and data."""
        self._setup_attendance_and_payroll(finalize=True)
        url = f"/api/payroll/periods/2026/8/export-bank-file/"

        expected_columns = [
            'Employee Code',
            'Beneficiary Name',
            'Account Number',
            'IFSC',
            'Net Amount',
            'Mode',
            'Remarks',
        ]

        res_csv = self.client.post(url, {"template": "GENERIC_NEFT", "format": "CSV"})
        self.assertEqual(res_csv.status_code, status.HTTP_200_OK)

        csv_content = res_csv.content.decode('utf-8-sig')
        reader = list(csv.reader(io.StringIO(csv_content)))
        self.assertEqual(reader[0], expected_columns)
        self.assertEqual(reader[1][0], "EMP-0001")
        self.assertEqual(reader[1][1], "Alice Smith")
        self.assertEqual(reader[1][2], "50100123456789")
        self.assertEqual(reader[1][3], "HDFC0000001")
        self.assertEqual(reader[1][4], "50000.00")

    def test_05_icici_and_sbi_templates(self):
        """ICICI and SBI templates format correctly and use corporate debit account."""
        self._setup_attendance_and_payroll(finalize=True)
        url = f"/api/payroll/periods/2026/8/export-bank-file/"

        # ICICI check
        res_icici = self.client.post(url, {"template": "ICICI", "format": "CSV"})
        self.assertEqual(res_icici.status_code, status.HTTP_200_OK)
        reader_icici = list(csv.reader(io.StringIO(res_icici.content.decode('utf-8-sig'))))
        self.assertEqual(reader_icici[0], [
            'Debit Account',
            'Beneficiary Account',
            'Beneficiary Name',
            'IFSC',
            'Amount',
            'Narration',
        ])
        # Corporate debit account from OrgSettings
        self.assertEqual(reader_icici[1][0], "998877665544")
        self.assertEqual(reader_icici[1][1], "50100123456789")
        self.assertEqual(reader_icici[1][2], "Alice Smith")

        # SBI check
        res_sbi = self.client.post(url, {"template": "SBI", "format": "CSV"})
        self.assertEqual(res_sbi.status_code, status.HTTP_200_OK)
        reader_sbi = list(csv.reader(io.StringIO(res_sbi.content.decode('utf-8-sig'))))
        self.assertEqual(reader_sbi[0], [
            'Beneficiary Account',
            'Amount',
            'Beneficiary Name',
            'IFSC',
            'Remarks',
        ])
        self.assertEqual(reader_sbi[1][0], "50100123456789")
        self.assertEqual(reader_sbi[1][1], "50000.00")
        self.assertEqual(reader_sbi[1][2], "Alice Smith")

    def test_06_tenant_isolation(self):
        """Org B admin cannot access or export Org A's payroll period."""
        self._setup_attendance_and_payroll(finalize=True)
        url = f"/api/payroll/periods/2026/8/export-bank-file/"

        # Authenticate as Admin B (different organization)
        self.client.force_authenticate(user=self.admin_b)

        res_get = self.client.get(url)
        self.assertEqual(res_get.status_code, status.HTTP_404_NOT_FOUND)

        res_post = self.client.post(url, {"template": "HDFC", "format": "CSV"})
        self.assertEqual(res_post.status_code, status.HTTP_404_NOT_FOUND)

    def test_07_unauthorized_access_denied(self):
        """User without payroll permissions cannot export bank files."""
        self._setup_attendance_and_payroll(finalize=True)
        url = f"/api/payroll/periods/2026/8/export-bank-file/"

        self.client.force_authenticate(user=self.viewer_a)

        res_get = self.client.get(url)
        self.assertEqual(res_get.status_code, status.HTTP_403_FORBIDDEN)

        res_post = self.client.post(url, {"template": "HDFC", "format": "CSV"})
        self.assertEqual(res_post.status_code, status.HTTP_403_FORBIDDEN)
