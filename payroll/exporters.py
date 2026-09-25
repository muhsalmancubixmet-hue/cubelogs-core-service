"""
Bank Payment Exporter Engine
Generates bank-compatible salary payment disbursement files (Excel/CSV)
from finalized payroll periods.
Supports HDFC, ICICI, SBI, and GENERIC_NEFT formats.
"""

from decimal import Decimal
import io
import csv
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


class BankPaymentExporter:
    """
    Modular salary payment exporter supporting multiple bank templates (NEFT/RTGS/CMS).
    """

    TEMPLATES = {
        'HDFC': {
            'name': 'HDFC Bank NEFT/RTGS',
            'columns': [
                'Beneficiary Name',
                'Account Number',
                'IFSC Code',
                'Amount',
                'Payment Mode',
                'Remarks',
            ],
        },
        'ICICI': {
            'name': 'ICICI Bank CMS/NEFT',
            'columns': [
                'Debit Account',
                'Beneficiary Account',
                'Beneficiary Name',
                'IFSC',
                'Amount',
                'Narration',
            ],
        },
        'SBI': {
            'name': 'State Bank of India CMP/NEFT',
            'columns': [
                'Beneficiary Account',
                'Amount',
                'Beneficiary Name',
                'IFSC',
                'Remarks',
            ],
        },
        'GENERIC_NEFT': {
            'name': 'Generic NEFT / Bulk Salary Upload',
            'columns': [
                'Employee Code',
                'Beneficiary Name',
                'Account Number',
                'IFSC',
                'Net Amount',
                'Mode',
                'Remarks',
            ],
        },
    }

    def __init__(self, payroll_period, template='GENERIC_NEFT', debit_account=None, remarks=None):
        self.period = payroll_period
        self.organization = payroll_period.organization
        template_upper = (template or '').upper()
        self.template_key = template_upper if template_upper in self.TEMPLATES else 'GENERIC_NEFT'
        self.template_info = self.TEMPLATES[self.template_key]

        # Resolve debit account
        if debit_account:
            self.debit_account = str(debit_account).strip()
        else:
            org_settings = getattr(self.organization, 'settings', None)
            self.debit_account = getattr(org_settings, 'corporate_account_number', '') or ''

        default_remarks = f"Salary {self.period.month:02d}/{self.period.year}"
        self.remarks = remarks.strip() if remarks else default_remarks

        self._prepared = False
        self.payable_records = []
        self.unpayable_employees = []
        self.total_payable_amount = Decimal('0.00')

    def prepare_data(self):
        """
        Fetches current PayrollEmployeeSnapshot records where net_payable > 0,
        validates bank details, and categorizes into payable and unpayable lists.
        """
        if self._prepared:
            return

        from payroll.models import PayrollEmployeeSnapshot

        snapshots = (
            PayrollEmployeeSnapshot.objects.filter(
                payroll_period=self.period,
                is_current=True,
                is_deleted=False,
                net_payable__gt=0,
            )
            .select_related('employee', 'employee__employee_profile')
            .order_by('employee__first_name', 'employee__last_name', 'id')
        )

        self.payable_records = []
        self.unpayable_employees = []
        self.total_payable_amount = Decimal('0.00')

        for snap in snapshots:
            emp = snap.employee
            prof = getattr(emp, 'employee_profile', None)

            # Resolve bank details (check Employee, then EmployeeProfile fallback)
            acc_num = getattr(emp, 'account_number', None) or (getattr(prof, 'account_number', None) if prof else None) or ''
            ifsc = getattr(emp, 'ifsc_code', None) or (getattr(prof, 'ifsc_code', None) if prof else None) or ''
            holder_name = getattr(emp, 'account_holder_name', None) or (getattr(prof, 'account_holder_name', None) if prof else None)

            if not holder_name:
                if emp and hasattr(emp, 'get_full_name') and emp.get_full_name().strip():
                    holder_name = emp.get_full_name().strip()
                elif snap.employee_name:
                    holder_name = snap.employee_name.strip()
                else:
                    holder_name = getattr(emp, 'email', '')

            emp_code = getattr(emp, 'employee_code', None) or (getattr(prof, 'employee_code', None) if prof else None) or (f"EMP-{emp.id:04d}" if emp else f"EMP-{snap.id:04d}")
            bank_name = getattr(emp, 'bank_name', None) or (getattr(prof, 'bank_name', None) if prof else None) or ''

            acc_clean = str(acc_num).strip().replace(' ', '').replace('-', '')
            ifsc_clean = str(ifsc).strip().upper()

            missing = []
            if not acc_clean:
                missing.append('account_number')
            if not ifsc_clean:
                missing.append('ifsc_code')

            amount = Decimal(str(snap.net_payable)).quantize(Decimal('0.01'))

            if missing:
                self.unpayable_employees.append({
                    'employee_id': emp.id if emp else None,
                    'snapshot_id': snap.id,
                    'employee_name': snap.employee_name,
                    'employee_code': emp_code,
                    'net_payable': str(amount),
                    'missing_fields': missing,
                    'bank_name': bank_name,
                })
            else:
                mode = 'RTGS' if amount >= Decimal('200000.00') else 'NEFT'
                self.payable_records.append({
                    'employee_id': emp.id if emp else None,
                    'snapshot_id': snap.id,
                    'employee_code': emp_code,
                    'beneficiary_name': holder_name,
                    'account_number': acc_clean,
                    'ifsc_code': ifsc_clean,
                    'amount': amount,
                    'payment_mode': mode,
                    'remarks': self.remarks,
                    'debit_account': self.debit_account,
                })
                self.total_payable_amount += amount

        self._prepared = True

    def get_summary(self):
        """
        Returns JSON-serializable pre-check summary.
        """
        self.prepare_data()
        return {
            'period_year': self.period.year,
            'period_month': self.period.month,
            'period_status': self.period.status,
            'template': self.template_key,
            'template_name': self.template_info['name'],
            'columns': self.template_info['columns'],
            'debit_account': self.debit_account,
            'currency': self.period.currency,
            'total_payable_count': len(self.payable_records),
            'total_payable_amount': str(self.total_payable_amount),
            'unpayable_count': len(self.unpayable_employees),
            'unpayable_employees': self.unpayable_employees,
            'remarks': self.remarks,
        }

    def _build_row(self, rec):
        """
        Transforms a record dictionary into the list of column values according to active template.
        """
        t = self.template_key
        if t == 'HDFC':
            return [
                rec['beneficiary_name'],
                rec['account_number'],
                rec['ifsc_code'],
                rec['amount'],
                rec['payment_mode'],
                rec['remarks'],
            ]
        elif t == 'ICICI':
            return [
                rec['debit_account'],
                rec['account_number'],
                rec['beneficiary_name'],
                rec['ifsc_code'],
                rec['amount'],
                rec['remarks'],
            ]
        elif t == 'SBI':
            return [
                rec['account_number'],
                rec['amount'],
                rec['beneficiary_name'],
                rec['ifsc_code'],
                rec['remarks'],
            ]
        else:  # GENERIC_NEFT
            return [
                rec['employee_code'],
                rec['beneficiary_name'],
                rec['account_number'],
                rec['ifsc_code'],
                rec['amount'],
                rec['payment_mode'],
                rec['remarks'],
            ]

    def generate_file(self, file_format='XLSX'):
        """
        Generates and returns (buffer, filename, content_type).
        """
        self.prepare_data()
        fmt = (file_format or 'XLSX').upper()

        base_name = f"Bank_Payment_{self.template_key}_{self.period.year}_{self.period.month:02d}"

        if fmt == 'CSV':
            buffer = io.StringIO()
            writer = csv.writer(buffer)
            # Write header
            writer.writerow(self.template_info['columns'])
            for rec in self.payable_records:
                row = self._build_row(rec)
                formatted_row = [
                    f"{val:.2f}" if isinstance(val, (Decimal, float)) else str(val)
                    for val in row
                ]
                writer.writerow(formatted_row)

            bytes_buffer = io.BytesIO(buffer.getvalue().encode('utf-8-sig'))
            return bytes_buffer, f"{base_name}.csv", 'text/csv; charset=utf-8'

        else:  # XLSX
            wb = Workbook()
            ws = wb.active
            ws.title = "Payment_Export"
            ws.views.sheetView[0].showGridLines = True

            header_font = Font(name='Segoe UI', size=11, bold=True, color='FFFFFF')
            header_fill = PatternFill(start_color='1E3A8A', end_color='1E3A8A', fill_type='solid')
            header_align = Alignment(horizontal='center', vertical='center', wrap_text=True)

            data_font = Font(name='Segoe UI', size=10)
            text_align = Alignment(vertical='center')
            num_align = Alignment(horizontal='right', vertical='center')
            center_align = Alignment(horizontal='center', vertical='center')

            thin_border = Border(
                left=Side(style='thin', color='CBD5E1'),
                right=Side(style='thin', color='CBD5E1'),
                top=Side(style='thin', color='CBD5E1'),
                bottom=Side(style='thin', color='CBD5E1'),
            )

            # Write headers
            headers = self.template_info['columns']
            ws.append(headers)
            ws.row_dimensions[1].height = 26

            for col_idx in range(1, len(headers) + 1):
                cell = ws.cell(row=1, column=col_idx)
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = header_align

            # Write rows
            for row_idx, rec in enumerate(self.payable_records, start=2):
                row_data = self._build_row(rec)
                ws.row_dimensions[row_idx].height = 20

                for col_idx, val in enumerate(row_data, start=1):
                    cell = ws.cell(row=row_idx, column=col_idx)
                    cell.font = data_font
                    cell.border = thin_border

                    if isinstance(val, (Decimal, float)):
                        cell.value = float(val)
                        cell.number_format = '0.00'
                        cell.alignment = num_align
                    else:
                        str_val = str(val) if val is not None else ""
                        cell.value = str_val
                        # Enforce explicit text type for account numbers and codes to avoid scientific notation
                        cell.data_type = 's'
                        if col_idx in [2, 3] and self.template_key in ['HDFC', 'GENERIC_NEFT']:
                            cell.alignment = center_align
                        elif self.template_key in ['ICICI', 'SBI'] and col_idx in [1, 2, 4]:
                            cell.alignment = center_align
                        else:
                            cell.alignment = text_align

            # Auto-fit column widths
            for col in ws.columns:
                max_len = 0
                col_letter = get_column_letter(col[0].column)
                for cell in col:
                    val_str = str(cell.value or '')
                    if len(val_str) > max_len:
                        max_len = len(val_str)
                ws.column_dimensions[col_letter].width = max(max_len + 4, 14)

            buffer = io.BytesIO()
            wb.save(buffer)
            buffer.seek(0)
            return buffer, f"{base_name}.xlsx", 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
